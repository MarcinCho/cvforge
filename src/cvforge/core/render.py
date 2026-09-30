"""TailoredCV + template -> PDF (WeasyPrint) and DOCX (python-docx).

A template is a folder with template.toml + style.css and optionally its own cv.html.j2
(otherwise the shared one is used). The DOCX look comes from template.toml [docx].
"""

from __future__ import annotations

import re
import shutil
import tomllib
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from jinja2 import ChoiceLoader, Environment, FileSystemLoader, select_autoescape
from pydantic import BaseModel, ConfigDict, Field

from ..config import BUILTIN_TEMPLATES, PACKAGE_DIR
from ..models import TailoredCV

FONTS_DIR = PACKAGE_DIR / "builtin" / "fonts"
SHARED_DIR = BUILTIN_TEMPLATES / "_shared"
FONT_FILES = {
    "NotoSans": (
        "Noto Sans",
        {
            "400 normal": "NotoSans-Regular.ttf",
            "700 normal": "NotoSans-Bold.ttf",
            "600 normal": "NotoSans-SemiBold.ttf",
            "400 italic": "NotoSans-Italic.ttf",
        },
    ),
    "Lato": (
        "Lato",
        {
            "400 normal": "Lato-Regular.ttf",
            "700 normal": "Lato-Bold.ttf",
            "600 normal": "Lato-Semibold.ttf",
            "400 italic": "Lato-Italic.ttf",
        },
    ),
    "NotoSerif": (
        "Noto Serif",
        {
            "400 normal": "NotoSerif-Regular.ttf",
            "700 normal": "NotoSerif-Bold.ttf",
            "600 normal": "NotoSerif-SemiBold.ttf",
            "400 italic": "NotoSerif-Italic.ttf",
        },
    ),
    "NotoSansMono": (
        "Noto Sans Mono",
        {"400 normal": "NotoSansMono-Regular.ttf", "700 normal": "NotoSansMono-Bold.ttf"},
    ),
    "Carlito": (
        "Carlito",
        {"400 normal": "Carlito-Regular.ttf", "700 normal": "Carlito-Bold.ttf", "400 italic": "Carlito-Italic.ttf"},
    ),
    "Caladea": (
        "Caladea",
        {"400 normal": "Caladea-Regular.ttf", "700 normal": "Caladea-Bold.ttf", "400 italic": "Caladea-Italic.ttf"},
    ),
}
# small UI words that are not section titles (e.g. a sidebar "Contact" heading)
LABELS = {
    "contact": {
        "pl": "Kontakt",
        "en": "Contact",
        "de": "Kontakt",
        "no": "Kontakt",
        "nb": "Kontakt",
        "sv": "Kontakt",
        "da": "Kontakt",
        "cs": "Kontakt",
        "fr": "Contact",
        "es": "Contacto",
    },
}
PRESENT = {
    "pl": "obecnie",
    "en": "Present",
    "de": "heute",
    "no": "nå",
    "nb": "nå",
    "sv": "nu",
    "da": "nu",
    "cs": "dosud",
    "fr": "aujourd'hui",
    "es": "actualidad",
}
DEFAULT_SECTIONS = ["summary", "experience", "skills", "projects", "education", "certifications", "languages"]
MIN_FONT_PT = 7.0


class DocxStyle(BaseModel):
    font: str = "Calibri"
    body_size: float = 10.5
    name_size: float = 20
    heading_size: float = 12
    accent: str = "000000"
    muted: str = "444444"
    margins_cm: float = 1.8
    heading_rule: bool = True


class FontFace(BaseModel):
    family: str
    file: str = Field(description="Path relative to the template folder, or a bundled font file name")
    weight: int = 400
    style: str = "normal"


class Template(BaseModel):
    """template.toml. Unknown keys are kept and available in cv.html.j2 as `template.<key>`."""

    model_config = ConfigDict(extra="allow")

    name: str
    description: str = ""
    ats_safe: bool = True
    max_pages: int = 2
    skill_tags: bool = False
    sections: list[str] = Field(default_factory=lambda: list(DEFAULT_SECTIONS))
    fonts: list[str] = Field(default_factory=lambda: ["NotoSans"], description="Bundled font families")
    font: list[FontFace] = Field(default_factory=list, description="Extra [[font]] faces, e.g. template-local files")
    docx: DocxStyle = Field(default_factory=DocxStyle)
    path: Path


class LayoutReport(BaseModel):
    pages: int
    overflow: list[str] = Field(default_factory=list)
    min_font_pt: float = 0
    last_page_chars: int = 0
    density: int = 0


# ------------------------------------------------------------------ template discovery


def list_templates(dirs: list[Path]) -> dict[str, Template]:
    out: dict[str, Template] = {}
    for d in reversed(dirs):  # user dirs override built-ins
        if not d.is_dir():
            continue
        for toml in sorted(d.glob("*/template.toml")):
            if toml.parent.name.startswith("_"):
                continue
            data = tomllib.loads(toml.read_text(encoding="utf-8"))
            data.setdefault("name", toml.parent.name)
            out[data["name"]] = Template(**data, path=toml.parent)
    return out


def get_template(name: str, dirs: list[Path]) -> Template:
    templates = list_templates(dirs)
    if name not in templates:
        raise KeyError(f"unknown template {name!r}; available: {', '.join(templates)}")
    return templates[name]


def new_template(name: str, base: Template, target_dir: Path) -> Path:
    if not re.fullmatch(r"[a-z0-9][a-z0-9_-]*", name):
        raise ValueError("template name: lowercase letters, digits, - and _ only")
    dest = target_dir / name
    if dest.exists():
        raise FileExistsError(f"{dest} already exists")
    shutil.copytree(base.path, dest)
    if not (dest / "cv.html.j2").exists():
        shutil.copy(SHARED_DIR / "cv.html.j2", dest / "cv.html.j2")
    toml = dest / "template.toml"
    toml.write_text(
        re.sub(r"^name\s*=.*$", f'name = "{name}"', toml.read_text(encoding="utf-8"), count=1, flags=re.M),
        encoding="utf-8",
    )
    return dest


# ------------------------------------------------------------------ HTML / PDF


def _date(value: str, lang: str) -> str:
    if not value:
        return ""
    if value.strip().lower() in ("present", "now", "current", "obecnie", "nå", "heute", "dziś", "teraz"):
        return PRESENT.get(lang, "Present")
    return value


def _display_url(url: str) -> str:
    return re.sub(r"^https?://(www\.)?", "", url).rstrip("/")


def _fonts_css(tpl: Template) -> str:
    rules = []
    for face in tpl.font:
        path = tpl.path / face.file
        if not path.exists():
            path = FONTS_DIR / face.file
        if not path.exists():
            raise FileNotFoundError(f"template {tpl.name!r}: font file {face.file!r} not found")
        rules.append(
            f'@font-face {{ font-family: "{face.family}"; font-weight: {face.weight}; '
            f'font-style: {face.style}; src: url("{path.resolve().as_uri()}"); }}'
        )
    for key in tpl.fonts:
        if key not in FONT_FILES:
            continue
        family, files = FONT_FILES[key]
        for spec, fname in files.items():
            weight, style = spec.split()
            rules.append(
                f'@font-face {{ font-family: "{family}"; font-weight: {weight}; font-style: {style}; '
                f'src: url("{(FONTS_DIR / fname).as_uri()}"); }}'
            )
    return "\n".join(rules)


def render_html(cv: TailoredCV, tpl: Template, density: int = 0) -> str:
    env = Environment(
        loader=ChoiceLoader([FileSystemLoader(tpl.path), FileSystemLoader(SHARED_DIR)]),
        autoescape=select_autoescape(["html", "j2"]),
    )
    env.filters["date"] = lambda v: _date(v, cv.language)
    env.filters["display_url"] = _display_url
    env.globals["label"] = lambda key: LABELS.get(key, {}).get(cv.language, LABELS.get(key, {}).get("en", key))
    css = (
        (SHARED_DIR / "base.css").read_text(encoding="utf-8")
        + "\n"
        + (tpl.path / "style.css").read_text(encoding="utf-8")
    )
    return env.get_template("cv.html.j2").render(
        cv=cv, template=tpl, sections=tpl.sections, css=css, fonts_css=_fonts_css(tpl), density=density
    )


def _walk(box):
    yield box
    for child in getattr(box, "children", []) or []:
        yield from _walk(child)


MAX_DENSITY = 2
NEAR_EMPTY_CHARS = 250


def _needs_fit(r: LayoutReport, max_pages: int) -> bool:
    return r.pages > max_pages or (r.pages > 1 and r.last_page_chars < NEAR_EMPTY_CHARS)


def render_pdf(cv: TailoredCV, tpl: Template, out: Path, autofit: bool = True) -> LayoutReport:
    """Render; if the last page is nearly empty or the page target is exceeded, tighten spacing a little."""
    best = None
    max_density = MAX_DENSITY + (1 if tpl.max_pages == 1 else 0)  # one-pagers may go a notch tighter
    for density in range(max_density + 1 if autofit else 1):
        html = render_html(cv, tpl, density)
        doc, report = _layout(html, tpl)
        report.density = density
        if best is None or (report.pages < best[1].pages):
            best = (html, report, doc)
        if not _needs_fit(report, tpl.max_pages):
            best = (html, report, doc)
            break
    html, report, doc = best
    out.with_suffix(".html").write_text(html, encoding="utf-8")
    doc.write_pdf(out)
    return report


def _removals(cv: TailoredCV):
    """Remove content ONE item at a time, least important first, yielding a label per removal.
    Bullets are ordered by relevance by the tailoring step, so cutting from the end keeps the best ones.
    Nothing is ever rewritten or added."""
    exp = cv.experience

    def trim_role(i: int, keep: int):
        while len(exp) > i and len(exp[i].bullets) > keep:
            exp[i].bullets.pop()
            yield f"{exp[i].title} ({exp[i].company}): bullet"

    def trim_older(keep: int):
        for i in range(len(exp) - 1, 1, -1):  # oldest role first
            yield from trim_role(i, keep)

    while cv.projects:
        if len(cv.projects[-1].bullets) > 1:
            cv.projects[-1].bullets.pop()
            yield "project bullet"
        else:
            cv.projects.pop()
            yield "project"
    yield from trim_older(1)
    yield from trim_role(1, 3)
    yield from trim_role(0, 6)
    while len(cv.skills) > 4:
        cv.skills.pop()
        yield "skill group"
    for g in cv.skills:
        while len(g.items) > 8:
            g.items.pop()
            yield "skill"
    for e in cv.education:
        if e.notes:
            e.notes = ""
            yield "education note"
    while len(cv.certifications) > 2:
        cv.certifications.pop()
        yield "certification"
    yield from trim_role(0, 5)
    yield from trim_role(1, 2)
    yield from trim_older(0)
    sentences = re.split(r"(?<=[.!?])\s+", cv.summary.strip())
    while len(sentences) > 2:
        sentences.pop()
        cv.summary = " ".join(sentences)
        yield "summary sentence"
    yield from trim_role(0, 4)
    yield from trim_role(1, 1)
    yield from trim_role(0, 3)


def fit_to_pages(cv: TailoredCV, tpl: Template, out: Path) -> tuple[TailoredCV, LayoutReport, list[str]]:
    """Fit into tpl.max_pages: first tighter spacing, then remove the least important content one item
    at a time until it fits (never adds or rewrites text). Returns the fitted CV and what was removed."""
    from collections import Counter

    current = cv.model_copy(deep=True)
    tightest = MAX_DENSITY + (1 if tpl.max_pages == 1 else 0)
    removed: Counter[str] = Counter()
    fits = _layout(render_html(current, tpl, tightest), tpl)[1].pages <= tpl.max_pages
    if not fits:
        for label in _removals(current):
            removed[label] += 1
            if _layout(render_html(current, tpl, tightest), tpl)[1].pages <= tpl.max_pages:
                break
    report = render_pdf(current, tpl, out)  # final render picks the loosest spacing that fits
    return current, report, [f"{n}× {label}" if n > 1 else label for label, n in removed.items()]


def _layout(html: str, tpl: Template):
    from weasyprint import HTML  # heavy import, keep lazy

    doc = HTML(string=html, base_url=str(tpl.path)).render()
    report = LayoutReport(pages=len(doc.pages))
    min_px = 1e9
    for n, page in enumerate(doc.pages, 1):
        root = page._page_box
        try:
            right_edge = root.content_box_x() + root.width
        except Exception:
            right_edge = page.width
        chars = 0
        for box in _walk(root):
            text = getattr(box, "text", None)
            if text is None:
                continue
            chars += len(text.strip())
            size = box.style["font_size"]
            if text.strip():
                min_px = min(min_px, size)
            if box.position_x + box.width > right_edge + 0.5:
                report.overflow.append(f"page {n}: text runs past the page edge: {text.strip()[:60]!r}")
        if n == len(doc.pages):
            report.last_page_chars = chars
    report.min_font_pt = round(min_px * 0.75, 1) if min_px < 1e9 else 0
    return doc, report


# ------------------------------------------------------------------ DOCX


def _rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_.upper())


def _bottom_border(paragraph, color: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", color)):
        bottom.set(qn(k), v)
    borders.append(bottom)
    p_pr.append(borders)


def _set_font(style, font: str, size: float | None = None) -> None:
    style.font.name = font
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        r_fonts.set(qn(attr), font)
    if size:
        style.font.size = Pt(size)


def render_docx(cv: TailoredCV, tpl: Template, out: Path) -> None:
    s = tpl.docx
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(s.margins_cm))
    usable = sec.page_width - sec.left_margin - sec.right_margin

    normal = doc.styles["Normal"]
    _set_font(normal, s.font, s.body_size)
    normal.paragraph_format.space_after = Pt(1)
    normal.paragraph_format.line_spacing = 1.08
    h1 = doc.styles["Heading 1"]
    _set_font(h1, s.font, s.heading_size)
    h1.font.bold, h1.font.color.rgb = True, _rgb(s.accent)
    h1.paragraph_format.space_before, h1.paragraph_format.space_after = Pt(9), Pt(3)
    h1.paragraph_format.keep_with_next = True
    bullet = doc.styles["List Bullet"]
    _set_font(bullet, s.font, s.body_size)
    bullet.paragraph_format.space_after = Pt(1)

    doc.core_properties.author = cv.contact.name
    doc.core_properties.title = f"{cv.contact.name} — {cv.headline}"
    doc.core_properties.keywords = ", ".join(cv.target_keywords)

    p = doc.add_paragraph()
    r = p.add_run(cv.contact.name)
    r.bold, r.font.size = True, Pt(s.name_size)
    p.paragraph_format.space_after = Pt(0)
    p = doc.add_paragraph()
    r = p.add_run(cv.headline)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(s.body_size * 1.2), _rgb(s.accent)
    parts = [cv.contact.location, cv.contact.phone, cv.contact.email, *map(_display_url, cv.contact.links)]
    p = doc.add_paragraph(" | ".join(x for x in parts if x))
    p.runs[0].font.color.rgb = _rgb(s.muted)

    def heading(text: str) -> None:
        h = doc.add_paragraph(text.upper(), style="Heading 1")
        if s.heading_rule:
            _bottom_border(h, s.accent)

    def entry(title: str, org: str, dates: str, meta: str = "") -> None:
        para = doc.add_paragraph()
        para.paragraph_format.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)
        para.paragraph_format.space_before = Pt(4)
        para.paragraph_format.keep_with_next = True
        run = para.add_run(title)
        run.bold = True
        if org:
            para.add_run(" | ")
            o = para.add_run(org)
            o.bold, o.font.color.rgb = True, _rgb(s.accent)
        if dates:
            d = para.add_run("\t" + dates)
            d.font.color.rgb = _rgb(s.muted)
        if meta:
            m = doc.add_paragraph()
            mr = m.add_run(meta)
            mr.italic, mr.font.color.rgb = True, _rgb(s.muted)
            m.paragraph_format.keep_with_next = True

    def span(a: str, b: str) -> str:
        a, b = _date(a, cv.language), _date(b, cv.language)
        return f"{a} – {b}" if a and b else a or b

    t = cv.section_titles
    for section in tpl.sections:
        if section == "summary" and cv.summary:
            heading(t.summary)
            doc.add_paragraph(cv.summary)
        elif section == "experience" and cv.experience:
            heading(t.experience)
            for job in cv.experience:
                entry(job.title, job.company, span(job.start, job.end), job.location)
                for b in job.bullets:
                    doc.add_paragraph(b.text, style="List Bullet")
        elif section == "projects" and cv.projects:
            heading(t.projects)
            for proj in cv.projects:
                entry(proj.name, "", "")
                for b in proj.bullets:
                    doc.add_paragraph(b.text, style="List Bullet")
        elif section == "skills" and cv.skills:
            heading(t.skills)
            for g in cv.skills:
                para = doc.add_paragraph()
                para.add_run(f"{g.category}: ").bold = True
                para.add_run(", ".join(g.items))
        elif section == "education" and cv.education:
            heading(t.education)
            for e in cv.education:
                entry(e.degree, e.institution, span(e.start, e.end), e.notes)
        elif section == "certifications" and cv.certifications:
            heading(t.certifications)
            for c in cv.certifications:
                doc.add_paragraph(c, style="List Bullet")
        elif section == "languages" and cv.languages:
            heading(t.languages)
            doc.add_paragraph(" · ".join(cv.languages))

    if cv.gdpr_clause:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        r = p.add_run(cv.gdpr_clause)
        r.font.size, r.font.color.rgb = Pt(7.5), _rgb(s.muted)
    doc.save(out)


# ------------------------------------------------------------------ template previews

SAMPLE_CV = PACKAGE_DIR / "builtin" / "sample_cv.json"


def _template_fingerprint(tpl: Template) -> str:
    import hashlib

    h = hashlib.sha1()
    for f in sorted([*tpl.path.rglob("*"), *SHARED_DIR.glob("*")]):
        if f.is_file():
            h.update(f"{f}:{f.stat().st_mtime_ns}".encode())
    return h.hexdigest()[:12]


def template_preview(tpl: Template, cache_dir: Path, cv: TailoredCV | None = None, scale: float = 0.9) -> Path:
    """First page of the template rendered with the sample CV, cached until the template changes."""
    import pypdfium2 as pdfium

    cv = cv or TailoredCV.model_validate_json(SAMPLE_CV.read_text(encoding="utf-8"))
    cache_dir.mkdir(parents=True, exist_ok=True)
    png = cache_dir / f"{tpl.name}-{_template_fingerprint(tpl)}.png"
    if png.exists():
        return png
    for old in cache_dir.glob(f"{tpl.name}-*.png"):
        old.unlink()
    pdf = cache_dir / f"{tpl.name}.pdf"
    render_pdf(cv, tpl, pdf)
    doc = pdfium.PdfDocument(str(pdf))
    doc[0].render(scale=scale).to_pil().save(png)
    doc.close()
    return png
