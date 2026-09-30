"""ATS + visual checks on rendered CV files."""

from __future__ import annotations

import re
import shutil
import subprocess
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from ..llm import LLMProvider, ask_json
from ..models import CheckReport, TailoredCV, VisualReview
from ..prompts import render_prompt
from .render import MIN_FONT_PT, LayoutReport, Template
from .text import fold

LIGATURES = re.compile("[ﬀ-ﬆ]")
KEYWORD_SURVIVAL_MIN = 0.9
EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
SPACED = re.compile(r"(?<!\w)(?:\w ){5,}\w(?!\w)")  # "P O D S U M O W A N I E": too much letter-spacing


def pdf_text(pdf: Path) -> str:
    """Text as a typical ATS sees it: poppler's pdftotext when available, else pypdf."""
    if shutil.which("pdftotext"):
        r = subprocess.run(["pdftotext", "-enc", "UTF-8", str(pdf), "-"], capture_output=True, text=True)
        if r.returncode == 0 and r.stdout.strip():
            return r.stdout
    return "\n".join(page.extract_text() or "" for page in PdfReader(str(pdf)).pages)


def _flat(s: str) -> str:
    # undo line-wrap hyphenation and whitespace differences
    return fold(re.sub(r"-\n(?=\w)", "", s))


def _content_checks(report: CheckReport, text: str, cv: TailoredCV) -> None:
    flat = _flat(text)

    if "�" in text or LIGATURES.search(text):
        report.add("text_encoding", "fail", "garbled characters or ligatures in extracted text")
    else:
        report.add("text_encoding", "pass", "extracted text is clean")

    spaced = SPACED.findall(text)
    if spaced:
        report.add(
            "letter_spacing",
            "fail",
            f"text extracts letter-by-letter ({spaced[0]!r}) — reduce CSS letter-spacing (keep it ≤ 0.05em)",
        )

    diacritics = sorted({c for c in cv.summary + cv.headline if ord(c) > 127 and c.isalpha()})
    if diacritics and not all(c in text for c in diacritics):
        missing = [c for c in diacritics if c not in text]
        report.add("diacritics", "fail", f"characters lost in extraction: {''.join(missing)}")
    elif diacritics:
        report.add("diacritics", "pass", f"diacritics preserved ({''.join(diacritics)})")

    c = cv.contact
    for label, value in (("name", c.name), ("email", c.email), ("phone", c.phone)):
        if not value:
            report.add(f"contact_{label}", "warn", f"no {label} in the CV")
        elif fold(value).replace(" ", "") in flat.replace(" ", ""):
            report.add(f"contact_{label}", "pass", value)
        else:
            report.add(f"contact_{label}", "fail", f"{label} {value!r} not found in extracted text")
    if c.email and not EMAIL.search(text):
        report.add("contact_email_parse", "fail", "no parseable e-mail address in extracted text")

    titles = cv.section_titles
    expected = [t for key, t in titles.model_dump().items() if _section_has_content(cv, key)]
    positions = [flat.find(fold(t)) for t in expected]
    missing = [t for t, p in zip(expected, positions, strict=True) if p < 0]
    if missing:
        report.add("section_headings", "fail", f"headings not found: {missing}")
    else:
        report.add("section_headings", "pass", f"{len(expected)} standard headings found")

    kws = [k for k in cv.target_keywords if k.strip()]
    if kws:
        found = [k for k in kws if fold(k) in flat]
        ratio = len(found) / len(kws)
        lost = [k for k in kws if k not in found]
        status = "pass" if ratio >= KEYWORD_SURVIVAL_MIN else "warn" if ratio >= 0.7 else "fail"
        report.add(
            "keywords_extractable",
            status,
            f"{len(found)}/{len(kws)} target keywords survive text extraction" + (f"; lost: {lost}" if lost else ""),
        )


def _section_has_content(cv: TailoredCV, key: str) -> bool:
    return bool(
        {
            "summary": cv.summary,
            "experience": cv.experience,
            "projects": cv.projects,
            "skills": cv.skills,
            "education": cv.education,
            "certifications": cv.certifications,
            "languages": cv.languages,
        }.get(key)
    )


def check_pdf(
    pdf: Path, cv: TailoredCV, tpl: Template, layout: LayoutReport | None = None, previews_dir: Path | None = None
) -> CheckReport:
    report = CheckReport(file=pdf.name)
    reader = PdfReader(str(pdf))
    report.pages = len(reader.pages)
    text = pdf_text(pdf)

    _content_checks(report, text, cv)

    # layout
    if report.pages <= tpl.max_pages:
        report.add("page_count", "pass", f"{report.pages} page(s), template target ≤ {tpl.max_pages}")
    else:
        status = "warn" if report.pages == tpl.max_pages + 1 else "fail"
        report.add(
            "page_count",
            status,
            f"{report.pages} pages, template target ≤ {tpl.max_pages} — shorten "
            "(try the 'concise' style or the 'compact' template)",
        )
    if layout:
        if layout.overflow:
            report.add("overflow", "fail", "; ".join(layout.overflow[:5]))
        else:
            report.add("overflow", "pass", "no text outside the printable area")
        if layout.min_font_pt and layout.min_font_pt < MIN_FONT_PT:
            report.add("min_font_size", "warn", f"smallest text is {layout.min_font_pt}pt")
        else:
            report.add("min_font_size", "pass", f"smallest text {layout.min_font_pt}pt")
        if report.pages > 1 and layout.last_page_chars < 250:
            report.add(
                "last_page_fill",
                "warn",
                f"last page is almost empty ({layout.last_page_chars} chars) — trim to save a page",
            )

    fonts = _pdf_fonts(reader)
    unembedded = [f for f, emb in fonts.items() if not emb]
    if unembedded:
        report.add("fonts_embedded", "fail", f"fonts not embedded: {unembedded}")
    else:
        report.add("fonts_embedded", "pass", ", ".join(sorted(fonts)) or "no fonts")

    if not tpl.ats_safe:
        report.add("template_ats_safe", "warn", f"template {tpl.name!r} is not marked ATS-safe")

    if previews_dir:
        report.previews = render_previews(pdf, previews_dir)
    return report


def _pdf_fonts(reader: PdfReader) -> dict[str, bool]:
    fonts: dict[str, bool] = {}
    for page in reader.pages:
        res = page.get("/Resources") or {}
        for _, ref in (res.get("/Font") or {}).items():
            f = ref.get_object()
            name = str(f.get("/BaseFont", "?")).split("+")[-1]
            desc = f.get("/FontDescriptor")
            if desc is None and "/DescendantFonts" in f:
                desc = f["/DescendantFonts"][0].get_object().get("/FontDescriptor")
            desc = desc.get_object() if desc is not None else {}
            fonts[name] = any(k in desc for k in ("/FontFile", "/FontFile2", "/FontFile3"))
    return fonts


def render_previews(pdf: Path, out_dir: Path, scale: float = 1.4) -> list[str]:
    import pypdfium2 as pdfium

    out_dir.mkdir(parents=True, exist_ok=True)
    for old in out_dir.glob(f"{pdf.stem}-page-*.png"):
        old.unlink()
    names = []
    doc = pdfium.PdfDocument(str(pdf))
    for i in range(len(doc)):
        img = doc[i].render(scale=scale).to_pil()
        name = f"{pdf.stem}-page-{i + 1}.png"
        img.save(out_dir / name)
        names.append(name)
    doc.close()
    return names


def check_docx(path: Path, cv: TailoredCV) -> CheckReport:
    report = CheckReport(file=path.name)
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)
    _content_checks(report, text, cv)

    if doc.tables:
        report.add("docx_tables", "warn", f"{len(doc.tables)} table(s) — some ATS skip table content")
    else:
        report.add("docx_tables", "pass", "no tables")
    body_xml = doc.element.body.xml
    if "txbxContent" in body_xml:
        report.add("docx_textboxes", "fail", "text boxes found — most ATS cannot read them")
    else:
        report.add("docx_textboxes", "pass", "no text boxes")
    header_text = " ".join(p.text for s in doc.sections for part in (s.header, s.footer) for p in part.paragraphs)
    if cv.contact.email and cv.contact.email in header_text and cv.contact.email not in text:
        report.add("docx_contact_in_body", "fail", "contact details only in header/footer")
    else:
        report.add("docx_contact_in_body", "pass", "contact details in document body")
    return report


def visual_review(previews: list[Path], provider: LLMProvider) -> VisualReview | None:
    if not getattr(provider, "supports_images", False) or not previews:
        return None
    return ask_json(provider, render_prompt("visual_review"), VisualReview, images=previews)


def merge_visual(report: CheckReport, review: VisualReview | None) -> None:
    if review is None:
        return
    status = {"ok": "pass", "minor_issues": "warn", "major_issues": "fail"}[review.verdict]
    report.add("visual_review", status, "; ".join(review.issues) or "no visual issues found")


def summarize(report: CheckReport) -> str:
    icons = {"pass": "✔", "warn": "!", "fail": "✘"}
    lines = [f"{report.file}: {report.status.upper()} ({report.pages or '-'} page(s))"]
    lines += [f"  {icons[i.status]} {i.rule}: {i.detail}" for i in report.items]
    return "\n".join(lines)
